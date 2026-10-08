"""Reads a Word document into blocks: headings (by style or outline level), paragraphs, lists and tables."""

from __future__ import annotations

import io
import re

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..model import Block, Document, Run

HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.I)


def style_outline_level(style) -> int | None:
    """Outline level from the style chain (custom heading styles set this instead of a 'Heading N' name)."""
    seen = 0
    while style is not None and seen < 10:
        m = HEADING_STYLE.match(style.name or "")
        if m:
            return int(m.group(1))
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            ol = ppr.find(qn("w:outlineLvl"))
            if ol is not None:
                lvl = int(ol.get(qn("w:val")))
                return lvl + 1 if lvl < 9 else None
        style = style.base_style
        seen += 1
    return None


def heading_level(p: Paragraph) -> int | None:
    ppr = p._p.pPr
    if ppr is not None:
        ol = ppr.find(qn("w:outlineLvl"))
        if ol is not None:
            lvl = int(ol.get(qn("w:val")))
            if lvl < 9:
                return lvl + 1
    try:
        return style_outline_level(p.style)
    except Exception:
        return None


def is_title(p: Paragraph) -> bool:
    try:
        return (p.style.name or "").lower() in ("title", "subtitle") and (p.style.name or "").lower() == "title"
    except Exception:
        return False


def para_runs(p: Paragraph) -> list[Run]:
    out: list[Run] = []
    style_bold = bool(p.style is not None and p.style.font is not None and p.style.font.bold)
    for r in p.runs:
        text = r.text
        if not text:
            continue
        run = Run(text, bold=bool(r.bold) or (r.bold is None and style_bold), italic=bool(r.italic), underline=bool(r.underline))
        if out and (out[-1].bold, out[-1].italic, out[-1].underline) == (run.bold, run.italic, run.underline):
            out[-1].text += text
        else:
            out.append(run)
    # Hyperlinks and fields hide text from p.runs; fall back to the full text when that happens.
    if "".join(r.text for r in out).strip() != p.text.strip() and p.text.strip():
        return [Run(p.text)]
    return out


def numbering_info(doc: DocxDocument, p: Paragraph) -> tuple[int, bool] | None:
    """(depth, ordered) when the paragraph is a list item."""
    ppr = p._p.pPr
    numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
    style_list = False
    if numpr is None:
        name = (p.style.name or "").lower() if p.style is not None else ""
        if name.startswith("list"):
            style_list = True
            depth = int(name[-1]) - 1 if name[-1].isdigit() else 0
            return depth, "number" in name
        return None
    ilvl_el = numpr.find(qn("w:ilvl"))
    numid_el = numpr.find(qn("w:numId"))
    depth = int(ilvl_el.get(qn("w:val"))) if ilvl_el is not None else 0
    if numid_el is None or numid_el.get(qn("w:val")) == "0":
        return None
    ordered = False
    try:
        numbering = doc.part.numbering_part.element
        num_id = numid_el.get(qn("w:val"))
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == num_id:
                abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
                for absn in numbering.findall(qn("w:abstractNum")):
                    if absn.get(qn("w:abstractNumId")) == abs_id:
                        for lvl in absn.findall(qn("w:lvl")):
                            if lvl.get(qn("w:ilvl")) == str(depth):
                                fmt = lvl.find(qn("w:numFmt"))
                                ordered = fmt is not None and fmt.get(qn("w:val")) not in ("bullet", "none")
    except Exception:
        pass
    _ = style_list
    return depth, ordered


def table_rows(t: Table) -> list[list[list[Run]]]:
    rows = []
    for row in t.rows:
        cells, seen = [], set()
        for c in row.cells:
            if id(c._tc) in seen:  # merged cells repeat
                continue
            seen.add(id(c._tc))
            runs: list[Run] = []
            for i, cp in enumerate(c.paragraphs):
                pr = para_runs(cp)
                if i and pr:
                    runs.append(Run("\n"))
                runs.extend(pr)
            cells.append(runs)
        rows.append(cells)
    return rows


def iter_body(doc: DocxDocument):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


_FAKE_HEADING = re.compile(r"^(?:(\d+(?:\.\d+)*)[.)]?\s+)?[A-Z][^.!?]{1,90}$")


def read_docx(data: bytes, filename: str) -> Document:
    d = docx.Document(io.BytesIO(data))
    blocks: list[Block] = []
    title = ""
    candidates: list[tuple[int, Paragraph]] = []  # bold short lines, for documents without heading styles
    for item in iter_body(d):
        if isinstance(item, Table):
            rows = table_rows(item)
            if any(any(r.text.strip() for c in row for r in c) for row in rows):
                blocks.append(Block("table", rows=rows))
            continue
        p = item
        text = p.text.strip()
        if not text:
            continue
        if is_title(p) and not title:
            title = text
            continue
        lvl = heading_level(p)
        if lvl:
            blocks.append(Block("heading", runs=[Run(text)], level=lvl))
            continue
        num = numbering_info(d, p)
        runs = para_runs(p)
        if num:
            blocks.append(Block("list", runs=runs, level=num[0], ordered=num[1]))
            continue
        blocks.append(Block("para", runs=runs))
        if runs and all(r.bold for r in runs if r.text.strip()) and _FAKE_HEADING.match(text) and len(text) < 90:
            candidates.append((len(blocks) - 1, p))

    notes: list[str] = []
    if not any(b.kind == "heading" for b in blocks) and candidates:
        # No heading styles: treat bold, short, sentence-less lines as headings ("1.2 Scope" -> level 2).
        for idx, p in candidates:
            m = _FAKE_HEADING.match(p.text.strip())
            depth = (m.group(1).count(".") + 1) if m and m.group(1) else 1
            blocks[idx] = Block("heading", runs=[Run(p.text.strip())], level=depth)
        notes.append("This document has no heading styles, so bold short lines were treated as headings.")

    title = title or (d.core_properties.title or "").strip()
    if not title:
        first_h = next((b for b in blocks if b.kind == "heading"), None)
        title = first_h.text if first_h and first_h.level == 1 and sum(1 for b in blocks if b.kind == "heading" and b.level == 1) == 1 else ""
    return Document(title=title or re.sub(r"\.[^.]+$", "", filename), blocks=blocks, kind="docx", notes=notes)
