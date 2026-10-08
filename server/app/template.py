"""Understands a Word template: its sections, which text is instructions to replace, and which fields to fill.

Conventions the template author can rely on:
  * Sections are paragraphs in Heading styles (or any style with an outline level).
  * Instruction text to be replaced is in italics, highlighted, grey, or wrapped in [brackets].
    Any other text under a heading is kept as fixed wording.
  * Fields are short placeholders: [Policy Owner], {{Effective Date}} or <<Version>>, anywhere
    (body, tables, headers, footers). They are filled from the source document or flagged.
"""

from __future__ import annotations

import io
import re
from dataclasses import asdict, dataclass, field

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .readers.docx_reader import heading_level, iter_body

FIELD = re.compile(r"\{\{\s*([^{}]{1,60}?)\s*\}\}|<<\s*([^<>]{1,60}?)\s*>>|\[\s*([^\[\]]{1,60}?)\s*\]")
INSTRUCTION_START = re.compile(
    r"^\W*(describe|explain|list|insert|enter|provide|include|state|identify|summari[sz]e|specify|outline|document|add|define|note|detail|replace|type)\b",
    re.I,
)
GREYS = {"808080", "7F7F7F", "A6A6A6", "999999", "666666", "595959", "BFBFBF", "8C8C8C", "767171", "757575"}


@dataclass
class TemplateSection:
    key: str
    title: str
    level: int
    parent: str | None
    path: list[str]
    guidance: str  # instruction text (replaced by source content)
    fixed: str  # wording kept as-is
    element_index: int  # body child index of the heading


@dataclass
class TemplateField:
    key: str
    name: str
    token: str  # the literal placeholder text, e.g. "[Policy Owner]"
    where: str  # body | header | footer | table


@dataclass
class TemplateInfo:
    title_placeholder: bool
    sections: list[TemplateSection] = field(default_factory=list)
    fields: list[TemplateField] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def to_json(self) -> dict:
        return asdict(self)


def field_name(m: re.Match) -> str:
    return (m.group(1) or m.group(2) or m.group(3) or "").strip()


def is_field_token(name: str) -> bool:
    words = name.split()
    return 0 < len(words) <= 6 and not INSTRUCTION_START.match(name) and not name.endswith((".", ":"))


def is_guidance(p: Paragraph) -> bool:
    text = p.text.strip()
    if not text:
        return False
    runs = [r for r in p.runs if r.text.strip()]
    if runs and all(bool(r.italic) for r in runs):
        return True
    if runs and all(r._r.find(qn("w:rPr")) is not None and r._r.rPr.find(qn("w:highlight")) is not None for r in runs):
        return True
    if runs and all(r.font.color is not None and r.font.color.rgb is not None and str(r.font.color.rgb).upper() in GREYS for r in runs):
        return True
    # A paragraph that is mostly one long bracketed instruction.
    for m in FIELD.finditer(text):
        name = field_name(m)
        if len(m.group(0)) >= 0.6 * len(text) and not is_field_token(name):
            return True
    if INSTRUCTION_START.match(text) and len(text.split()) >= 4 and (text.endswith(".") or text.endswith("]")) and len(text) < 400:
        return bool(re.search(r"\b(this section|here|the policy|this policy|this document|this procedure)\b", text, re.I))
    return False


def is_title_para(p: Paragraph) -> bool:
    try:
        return (p.style.name or "").lower() == "title"
    except Exception:
        return False


def _collect_fields(texts: list[tuple[str, str]]) -> list[TemplateField]:
    found: dict[str, TemplateField] = {}
    for text, where in texts:
        for m in FIELD.finditer(text):
            name = field_name(m)
            if is_field_token(name) and m.group(0) not in found:
                found[m.group(0)] = TemplateField(key=f"f{len(found) + 1}", name=name, token=m.group(0), where=where)
    return list(found.values())


def analyse_template(data: bytes) -> TemplateInfo:
    d = docx.Document(io.BytesIO(data))
    info = TemplateInfo(title_placeholder=False)
    stack: list[TemplateSection] = []
    current: TemplateSection | None = None
    field_texts: list[tuple[str, str]] = []

    for idx, item in enumerate(iter_body(d)):
        if isinstance(item, Table):
            for row in item.rows:
                for c in row.cells:
                    field_texts.append((c.text, "table"))
            if current:
                current.fixed += ("\n" if current.fixed else "") + "[table]"
            continue
        p = item
        text = p.text.strip()
        if is_title_para(p):
            info.title_placeholder = True
            field_texts.append((text, "body"))
            continue
        lvl = heading_level(p)
        if lvl and text:
            while stack and stack[-1].level >= lvl:
                stack.pop()
            current = TemplateSection(
                key=f"s{len(info.sections) + 1}",
                title=text,
                level=lvl,
                parent=stack[-1].key if stack else None,
                path=[s.title for s in stack],
                guidance="",
                fixed="",
                element_index=idx,
            )
            info.sections.append(current)
            stack.append(current)
            field_texts.append((text, "body"))
            continue
        if not text:
            continue
        if current and is_guidance(p):
            current.guidance += ("\n" if current.guidance else "") + text
        else:
            field_texts.append((text, "body"))
            if current:
                current.fixed += ("\n" if current.fixed else "") + text

    for s in d.sections:
        for part, where in ((s.header, "header"), (s.footer, "footer"), (s.first_page_header, "header"), (s.first_page_footer, "footer")):
            try:
                for p in part.paragraphs:
                    field_texts.append((p.text, where))
                for t in part.tables:
                    for row in t.rows:
                        for c in row.cells:
                            field_texts.append((c.text, where))
            except Exception:
                pass

    info.fields = _collect_fields(field_texts)
    if not info.sections:
        info.notes.append("No headings were found. Put each section title in a Heading style (Heading 1, Heading 2, ...) so documents can be matched to it.")
    if info.sections and not any(s.guidance for s in info.sections):
        info.notes.append("No instruction text was detected, so all wording under each heading is kept and source content is added after it. Put instructions in italics or [brackets] to have them replaced.")
    return info
