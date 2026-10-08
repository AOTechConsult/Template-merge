"""Sample template and source documents for demo mode and tests. All names and content are invented."""

from __future__ import annotations

import io
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

SAMPLES = Path(__file__).parent / "samples"


def _guidance(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    return p


def sample_template() -> bytes:
    d = docx.Document()
    d.core_properties.title = "Information Security Policy Template"
    hdr = d.sections[0].header.paragraphs[0]
    hdr.text = "[Policy Title]  |  Example Corp internal"
    d.sections[0].footer.paragraphs[0].text = "Classification: Internal. Uncontrolled when printed."
    d.add_paragraph("[Policy Title]", style="Title")
    t = d.add_table(rows=4, cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate([("Policy Owner", "[Policy Owner]"), ("Effective Date", "[Effective Date]"), ("Version", "[Version]"), ("Approved By", "[Approver]")]):
        t.cell(i, 0).text = k
        t.cell(i, 1).text = v
    d.add_heading("1. Purpose", 1)
    _guidance(d, "Describe why this policy exists and the risk it addresses.")
    d.add_heading("2. Scope", 1)
    d.add_paragraph("[Describe who and what this policy applies to: people, systems, data and locations.]")
    d.add_heading("3. Definitions", 1)
    _guidance(d, "List terms used in this policy and what they mean here.")
    d.add_heading("4. Roles and Responsibilities", 1)
    _guidance(d, "Identify who owns, approves, implements and follows this policy.")
    d.add_heading("5. Policy Statements", 1)
    _guidance(d, "State the requirements of this policy in the subsections below.")
    d.add_heading("5.1 Access Control", 2)
    _guidance(d, "Describe how access is requested, approved, reviewed and removed.")
    d.add_heading("5.2 Password Requirements", 2)
    _guidance(d, "Describe password and multi-factor authentication requirements.")
    d.add_heading("5.3 Monitoring and Logging", 2)
    _guidance(d, "Describe what activity is logged, where logs are kept, and how long.")
    d.add_heading("6. Exceptions", 1)
    d.add_paragraph("Exceptions must be requested through the IT Service Desk and approved by the CISO. Approved exceptions are reviewed annually.")
    d.add_heading("7. Enforcement", 1)
    d.add_paragraph("Violations of this policy may result in disciplinary action, up to and including termination of employment.")
    d.add_heading("8. Related Documents", 1)
    _guidance(d, "List related policies, standards and procedures.")
    d.add_heading("9. Revision History", 1)
    rt = d.add_table(rows=2, cols=4)
    rt.style = "Table Grid"
    for i, h in enumerate(["Version", "Date", "Author", "Changes"]):
        rt.cell(0, i).text = h
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def sample_source_docx() -> bytes:
    d = docx.Document()
    d.add_paragraph("Access Management Policy", style="Title")
    d.add_paragraph("Document Owner: Jordan Reyes")
    d.add_paragraph("Effective: March 1, 2026")
    d.add_paragraph("Version: 2.1")
    d.add_paragraph(
        "This policy establishes how access to company systems and data is granted, reviewed and removed, "
        "to reduce the risk of unauthorised access to customer and financial information."
    )
    d.add_heading("Who This Applies To", 1)
    d.add_paragraph("All employees, contractors and vendors with access to company systems, including cloud applications and on-premises servers.")
    d.add_heading("Requirements", 1)
    d.add_paragraph("The following requirements apply to every account.")
    d.add_heading("User Access Reviews", 2)
    d.add_paragraph("Managers review their team's access every quarter. Access no longer needed is removed within five business days.")
    d.add_paragraph("Requests are submitted through the IT Service Desk", style="List Bullet")
    d.add_paragraph("Access is approved by the data owner", style="List Bullet")
    d.add_paragraph("Terminated users are disabled the same day", style="List Bullet")
    d.add_heading("Passwords", 2)
    p = d.add_paragraph("Passwords must be at least ")
    p.add_run("14 characters").bold = True
    p.add_run(" and multi-factor authentication is required for all remote access.")
    d.add_heading("Privileged Accounts", 2)
    d.add_paragraph("Administrators use separate privileged accounts that are not used for email or browsing.")
    d.add_heading("Exception Process", 1)
    d.add_paragraph("Exceptions require a documented business reason and expire after twelve months.")
    d.add_heading("Contacts", 1)
    d.add_paragraph("Questions about this policy go to the IT Security team at security@example-corp.com.")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def sample_pdf() -> bytes:
    return (SAMPLES / "remote-work-standard.pdf").read_bytes()


CONFLUENCE_PAGE = (
    "Acceptable Use Policy",
    """
<ac:structured-macro ac:name="info"><ac:rich-text-body><p>Owner: IT Security. Version: 1.4</p></ac:rich-text-body></ac:structured-macro>
<h1>Introduction</h1>
<p>This policy sets out the acceptable use of company devices, networks and accounts.</p>
<h1>Applicability</h1>
<p>It applies to all <strong>employees</strong>, interns and contractors using company resources.</p>
<h1>Terms</h1>
<table><tbody>
<tr><th>Term</th><th>Meaning</th></tr>
<tr><td>Company resources</td><td>Devices, networks, accounts and data owned or managed by the company.</td></tr>
<tr><td>Personal use</td><td>Use not related to company business.</td></tr>
</tbody></table>
<h1>Acceptable Use Rules</h1>
<ul><li>Limited personal use is allowed if it does not affect work.</li>
<li>Do not install unapproved software.<ul><li>Request software through the Service Desk.</li></ul></li>
<li>Report lost or stolen devices within one hour.</li></ul>
<h2>Email and Messaging</h2>
<p>Do not forward company email to personal accounts.</p>
<h1>Violations</h1>
<p>Misuse is reported to HR and may lead to loss of access.</p>
""",
)


def demo_confluence_page(ref: str) -> tuple[str, str]:
    return CONFLUENCE_PAGE
