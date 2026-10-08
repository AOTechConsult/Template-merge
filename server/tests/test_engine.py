"""Engine tests on the bundled samples. Run from server/: DEMO_MODE=true python -m tests.test_engine"""

from __future__ import annotations

import io
import json
import os

os.environ.setdefault("DEMO_MODE", "true")
os.environ.setdefault("DATA_DIR", "/tmp/tmerge-test")

import docx  # noqa: E402
import httpx  # noqa: E402

from app import ai, demo  # noqa: E402
from app.config import settings  # noqa: E402
from app.engine import merge, read_source  # noqa: E402
from app.matcher import similarity  # noqa: E402
from app.template import analyse_template  # noqa: E402

TEMPLATE = demo.sample_template()
INFO = analyse_template(TEMPLATE)
key = {s.title: s.key for s in INFO.sections}


def text_of(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def test_template_analysis():
    assert len(INFO.sections) == 12
    assert [f.name for f in INFO.fields] == ["Policy Title", "Policy Owner", "Effective Date", "Version", "Approver"]
    assert INFO.sections[0].guidance and not INFO.sections[0].fixed  # italic instructions
    assert INFO.sections[1].guidance  # [bracketed] instructions
    assert INFO.sections[8].fixed and not INFO.sections[8].guidance  # fixed wording kept


def test_similarity():
    assert similarity("3.2 Roles & Responsibilities:", "Roles and Responsibilities") == 1.0
    assert similarity("Who This Applies To", "2. Scope") >= 0.85
    assert similarity("Passwords", "Password Requirements") >= 0.6
    assert similarity("Contacts", "Definitions") < 0.6


def test_docx_headings_only():
    r = merge(TEMPLATE, INFO, read_source("file", "a.docx", demo.sample_source_docx(), None), "a.docx", None)
    pre = {s["title"]: s for s in r.preview["sections"]}
    assert pre["2. Scope"]["placements"][0]["source_title"] == "Who This Applies To"
    assert pre["5.2 Password Requirements"]["placements"][0]["source_title"] == "Passwords"
    assert pre["5.1 Access Control"]["placements"][0]["source_title"] == "User Access Reviews"
    assert "3. Definitions" in r.report["missing_sections"]
    assert r.report["missing_fields"] == ["Approver"]
    out = text_of(r.output)
    assert "Jordan Reyes" in out and "[MISSING: Approver]" in out and "[Policy Owner]" not in out
    assert "Missing from the source document." in out
    assert "Content not in the template" in out  # preamble + Contacts kept, nothing dropped
    assert "14 characters" in out


def test_pdf_and_confluence():
    r = merge(TEMPLATE, INFO, read_source("file", "s.pdf", demo.sample_pdf(), None), "s.pdf", None)
    pre = {s["title"]: s for s in r.preview["sections"]}
    assert pre["1. Purpose"]["placements"][0]["source_title"] == "1. Objective"
    assert {f["name"]: f["value"] for f in r.preview["fields"]}["Policy Owner"] == "Priya Natarajan"
    r = merge(TEMPLATE, INFO, read_source("confluence", "c", None, "x"), "c", None)
    pre = {s["title"]: s for s in r.preview["sections"]}
    assert pre["3. Definitions"]["placements"][0]["blocks"][0]["kind"] == "table"
    assert pre["5. Policy Statements"]["placements"][0]["source_title"] == "Acceptable Use Rules"


def _fake_post(provider):
    """Simulates the provider's HTTP reply so request building and reply parsing are exercised."""
    calls = []

    def post(url, **kw):
        calls.append((url, kw))
        prompt = kw["json"]["messages"][-1]["content"]
        if "SOURCE PIECES" in prompt:
            ids = [ln.split(":")[0].strip("- ").strip() for ln in prompt.splitlines() if ln.startswith("- u")]
            body = {"assignments": [{"unit": i, "section": key["1. Purpose"] if n == 0 else None, "confidence": "high", "reason": "describes intent"} for n, i in enumerate(ids)]}
        else:
            body = {"fields": [{"field": INFO.fields[4].key, "value": "CISO", "quote": "approved by the CISO"}]}
        text = "```json\n" + json.dumps(body) + "\n```"
        if provider == "claude":
            payload = {"content": [{"type": "text", "text": text}]}
        else:
            payload = {"choices": [{"message": {"content": json.dumps(body)}}]}
        return httpx.Response(200, json=payload, request=httpx.Request("POST", url))

    return post, calls


def test_ai_providers():
    for provider in ("claude", "azure"):
        settings.claude.api_key = "test-key"
        settings.azure.endpoint, settings.azure.api_key, settings.azure.deployment = "https://example.openai.azure.com", "k", "gpt"
        post, calls = _fake_post(provider)
        original = ai.httpx.post
        ai.httpx.post = post
        try:
            r = merge(TEMPLATE, INFO, read_source("file", "a.docx", demo.sample_source_docx(), None), "a.docx", provider)
        finally:
            ai.httpx.post = original
        pre = {s["title"]: s for s in r.preview["sections"]}
        assert pre["1. Purpose"]["placements"][0]["method"] == "ai", provider
        assert r.report["ai"]["calls"] == 2 and r.report["missing_fields"] == []
        assert "Contacts" in [e["source_title"] for e in r.preview["extra"]]
        url = calls[0][0]
        assert ("anthropic.com" in url) if provider == "claude" else ("/openai/deployments/gpt/chat/completions" in url)
        assert "CISO" in text_of(r.output)


if __name__ == "__main__":
    n = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            n += 1
            print(f"ok  {name}")
    print(f"{n} tests passed")
